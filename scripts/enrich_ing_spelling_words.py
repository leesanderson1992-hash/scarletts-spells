#!/usr/bin/env python3
"""Enrich approved -ing word content with reusable dictionary evidence.

This produces a review candidate only.  User-supplied definitions and dictation
sentences are copied verbatim.  Lexical data and proposed transformations are
labelled as candidate evidence so they cannot be mistaken for an activated
teaching-dictionary release.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import statistics
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from wordfreq import zipf_frequency

ROOT = Path(__file__).resolve().parents[1]
INTAKE = ROOT / "docs/implementation/seed-data/teaching-dictionary/candidates/2026-06-29-phase-5-source-intake"
RELEASES = ROOT / "docs/implementation/seed-data/teaching-dictionary/releases"


def norm(value: Any) -> str:
    return str(value or "").strip().lower().replace("’", "'")


def word_key(word: str) -> str:
    return f"{re.sub(r'[^a-z0-9]+', '_', word).strip('_')}_en_gb"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def morpholex_index(path: Path) -> dict[str, dict[str, str]]:
    book = load_workbook(path, read_only=True, data_only=True)
    index: dict[str, dict[str, str]] = {}
    for sheet_name in book.sheetnames:
        if not re.fullmatch(r"\d+-\d+-\d+", sheet_name):
            continue
        rows = book[sheet_name].iter_rows(values_only=True)
        fields = [str(value or "") for value in next(rows)]
        if "Word" not in fields or "MorphoLexSegm" not in fields:
            continue
        positions = {field: i for i, field in enumerate(fields)}
        for values in rows:
            word = norm(values[positions["Word"]])
            if word and word not in index:
                index[word] = {
                    "segmentation": str(values[positions["MorphoLexSegm"]] or ""),
                    "pos": str(values[positions.get("POS", 0)] or ""),
                    "nmorph": str(values[positions.get("Nmorph", 0)] or ""),
                    "prs_signature": str(values[positions.get("PRS_signature", 0)] or ""),
                }
    return index


def ipa_index(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if "\t" not in line:
            continue
        word, ipa = line.split("\t", 1)
        result.setdefault(norm(word), ipa.strip())
    return result


def cmu_index(path: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    for line in path.read_text(encoding="latin-1").splitlines():
        if not line or line.startswith(";;;") or " " not in line:
            continue
        word, phones = line.split(" ", 1)
        result.setdefault(norm(re.sub(r"\(\d+\)$", "", word)), phones.strip())
    return result


def bnc_index() -> dict[str, float]:
    totals: dict[str, float] = defaultdict(float)
    path = INTAKE / "british_frequency_bnc_1_2_all_freq_source.txt"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            try:
                totals[norm(row.get("Word", "")).removesuffix("*")] += float(row.get("Freq", "") or 0)
            except ValueError:
                pass
    return totals


def aoa_index() -> dict[str, float]:
    path = INTAKE / "brysbaert_biemiller_test_based_aoa_master_source.xlsx"
    book = load_workbook(path, read_only=True, data_only=True)
    rows = book[book.sheetnames[0]].iter_rows(values_only=True)
    fields = [str(value or "") for value in next(rows)]
    word_i, age_i = fields.index("WORD"), fields.index("AoAtestbased")
    values: dict[str, list[float]] = defaultdict(list)
    for row in rows:
        word = norm(row[word_i])
        try:
            age = float(row[age_i])
        except (TypeError, ValueError):
            continue
        if word and " " not in word and "(" not in word:
            values[word].append(age)
    return {word: statistics.median(ages) for word, ages in values.items()}


def syllable_facts(ipa: str, cmu: str) -> tuple[str, str, str]:
    stresses = re.findall(r"[012]", cmu)
    if stresses:
        labels = {"0": "unstressed", "1": "primary", "2": "secondary"}
        return str(len(stresses)), "-".join(labels[value] for value in stresses), str("AH0" in cmu or "ER0" in cmu).upper()
    nuclei = re.findall(r"[aeiouæɑɒɔəɜɛɪʊʌɐ]+", ipa.replace("ː", ""))
    return (str(max(1, len(nuclei))) if ipa else "", "in_review" if ipa else "", str(any(symbol in ipa for symbol in ("ə", "ɐ"))).upper() if ipa else "")


def existing_canonical() -> dict[str, dict[str, str]]:
    """Latest release path wins and records the prior canonical projection."""
    current: dict[str, dict[str, str]] = {}
    for words_path in sorted(RELEASES.glob("*/package/canonical_words.csv")):
        package = words_path.parent
        metadata = {row["word_key"]: row for row in read_csv(package / "canonical_word_metadata.csv")} if (package / "canonical_word_metadata.csv").exists() else {}
        morphology = {row["word_key"]: row for row in read_csv(package / "canonical_word_morphology.csv")} if (package / "canonical_word_morphology.csv").exists() else {}
        dictation = {row["word_key"]: row for row in read_csv(package / "dictation_sentences.csv")} if (package / "dictation_sentences.csv").exists() else {}
        for word in read_csv(words_path):
            key = word["word_key"]
            current[word["normalised_word"]] = {
                "release": str(package.relative_to(ROOT)),
                "word_key": key,
                "review_status": word.get("review_status", ""),
                "row_status": word.get("row_status", ""),
                "existing_phoneme_hint": metadata.get(key, {}).get("phoneme_hint", ""),
                "existing_word_sum": morphology.get(key, {}).get("word_sum", ""),
                "existing_transformation_notes": morphology.get(key, {}).get("transformation_notes", ""),
                "existing_dictation_sentence": dictation.get(key, {}).get("dictation_sentence", ""),
            }
    return current


def base_and_rule(word: str, morph: dict[str, dict[str, str]], canonical: dict[str, dict[str, str]]) -> tuple[str, str]:
    """Propose one of the four target rules from source morphology and spelling."""
    stem = word[:-3] if word.endswith("ing") else ""
    existing_note = canonical.get(word, {}).get("existing_transformation_notes", "").casefold()
    existing_sum = canonical.get(word, {}).get("existing_word_sum", "")
    if existing_note and existing_sum:
        base = existing_sum.split("+", 1)[0].strip()
        if "double" in existing_note:
            return base, "double_final_consonant"
        if "final <e>" in existing_note or "drop the final e" in existing_note:
            return base, "drop_e"

    # A -ying ending is either a regular -y base or the ie → y transformation.
    # Prefer the regular route only where the raw lexical source identifies that
    # stem as a verb; this avoids treating abbreviations such as <ly> as bases.
    if word.endswith("ying"):
        ie_base = f"{word[:-4]}ie"
        if word == "retying":  # Approved source word; absent from this MorphoLex edition.
            return "retie", "ie_to_y"
        if "VB" not in morph.get(stem, {}).get("pos", ""):
            return ie_base, "ie_to_y"

    if len(stem) >= 2 and stem[-1:] == stem[-2:-1]:
        doubled_base = stem[:-1]
        return doubled_base, "double_final_consonant"

    drop_e_base = f"{stem}e"
    segmentation = morph.get(word, {}).get("segmentation", "")
    if f"({drop_e_base})" in segmentation:
        return drop_e_base, "drop_e"

    if stem:
        return stem, "regular"
    return "", "unresolved"

def transformed(base: str, rule: str) -> str:
    if rule == "regular": return f"{base}ing"
    if rule == "drop_e": return f"{base[:-1]}ing"
    if rule == "double_final_consonant": return f"{base}{base[-1]}ing"
    if rule == "ie_to_y": return f"{base[:-2]}ying"
    return ""


def transformation_note(base: str, rule: str) -> str:
    return {
        "regular": "Add -ing to the base word.",
        "drop_e": "Drop the final e, then add -ing.",
        "double_final_consonant": "Double the final consonant, then add -ing.",
        "ie_to_y": "Change ie to y, then add -ing.",
    }.get(rule, "")


def micro_skill(rule: str) -> str:
    return {
        "regular": "D4_INF_ING_ENDINGS_REGULAR",
        "drop_e": "D4_INF_ING_ENDINGS_DROP_E",
        "double_final_consonant": "D4_INF_ING_ENDINGS_DOUBLE_FINAL_CONSONANT",
        "ie_to_y": "D4_INF_ING_ENDINGS_IE_TO_Y",
    }.get(rule, "")


def doubling_pattern(base: str, rule: str) -> str:
    if rule != "double_final_consonant":
        return ""
    # British <travel → travelling> has a final-l convention, not the
    # stressed-final-syllable rule taught in this two-pattern lesson.
    if base == "travel":
        return "british_final_l"
    return "short_cvc" if len(base) <= 4 and len(re.findall(r"[aeiou]", base)) == 1 else "stressed_final_syllable"


def safe_for_v1(base: str, rule: str, pattern: str) -> str:
    if rule != "double_final_consonant" or pattern in {"short_cvc", "stressed_final_syllable"}:
        return "candidate"
    return "requires_british_final_l_rule"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--morpholex", type=Path, required=True)
    parser.add_argument("--ipa", type=Path, required=True)
    parser.add_argument("--cmudict", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    provided = read_csv(args.input)
    required = {"Word", "Child Friendly Definition", "Simple Dictation Sentence"}
    if not provided or required - set(provided[0]):
        raise SystemExit("Input must contain Word, Child Friendly Definition, and Simple Dictation Sentence.")
    words = [norm(row["Word"]) for row in provided]
    if len(words) != len(set(words)) or any(not re.fullmatch(r"[a-z]+", word) for word in words):
        raise SystemExit("Words must be unique lower-case alphabetic spellings.")

    morph = morpholex_index(args.morpholex)
    ipa, cmu, bnc, aoa, canonical = ipa_index(args.ipa), cmu_index(args.cmudict), bnc_index(), aoa_index(), existing_canonical()
    output_rows: list[dict[str, str]] = []
    for supplied in provided:
        word = norm(supplied["Word"])
        base, rule = base_and_rule(word, morph, canonical)
        valid = bool(base) and transformed(base, rule) == word
        pattern = doubling_pattern(base, rule)
        target_morph, base_morph = morph.get(word, {}), morph.get(base, {})
        target_ipa, target_cmu = ipa.get(word, ""), cmu.get(word, "")
        syllables, stress, schwa = syllable_facts(target_ipa, target_cmu)
        tokens = re.findall(r"[a-z]+(?:'[a-z]+)?", norm(supplied["Simple Dictation Sentence"]))
        output_rows.append({
            "word_key": word_key(word), "display_word": word, "normalised_word": word, "dialect_code": "en-GB",
            "provided_child_friendly_definition": supplied["Child Friendly Definition"].strip(),
            "provided_simple_dictation_sentence": supplied["Simple Dictation Sentence"].strip(),
            "provided_audio_text": supplied["Simple Dictation Sentence"].strip(),
            "provided_target_occurrences": str(tokens.count(word)),
            "provided_content_status": "approved_source_supplied_by_content_owner",
            "proposed_base_verb": base, "proposed_ing_rule": rule, "proposed_doubling_pattern": pattern, "proposed_micro_skill_key": micro_skill(rule),
            "proposed_word_sum": f"{base} + ing → {word}" if base else "",
            "proposed_transformation_note": transformation_note(base, rule),
            "spelling_transformation_status": "valid" if valid else "unresolved",
            "v1_lesson_eligibility": safe_for_v1(base, rule, pattern) if valid else "blocked",
            "target_morpholex_segmentation": target_morph.get("segmentation", ""),
            "target_morpholex_pos": target_morph.get("pos", ""), "target_morpholex_nmorph": target_morph.get("nmorph", ""),
            "target_morpholex_prs_signature": target_morph.get("prs_signature", ""),
            "base_morpholex_segmentation": base_morph.get("segmentation", ""), "base_morpholex_pos": base_morph.get("pos", ""),
            "target_british_ipa": target_ipa, "target_cmudict_phones": target_cmu,
            "target_syllables_candidate": syllables, "target_stress_pattern_candidate": stress, "target_has_schwa_candidate": schwa,
            "target_wordfreq_zipf": f"{zipf_frequency(word, 'en'):.3f}", "target_bnc_frequency": f"{bnc.get(word, 0):g}",
            "target_aoa_median": f"{aoa[word]:.2f}" if word in aoa else "",
            "existing_canonical_release": canonical.get(word, {}).get("release", ""),
            "existing_canonical_word_key": canonical.get(word, {}).get("word_key", ""),
            "existing_canonical_review_status": canonical.get(word, {}).get("review_status", ""),
            "existing_canonical_row_status": canonical.get(word, {}).get("row_status", ""),
            "existing_canonical_phoneme_hint": canonical.get(word, {}).get("existing_phoneme_hint", ""),
            "existing_canonical_word_sum": canonical.get(word, {}).get("existing_word_sum", ""),
            "existing_canonical_transformation_notes": canonical.get(word, {}).get("existing_transformation_notes", ""),
            "existing_canonical_dictation_sentence": canonical.get(word, {}).get("existing_dictation_sentence", ""),
            "morphology_review_status": "in_review", "lexical_sources": "MorphoLex-en | open-dict-data ipa-dict en_UK | CMUdict | wordfreq | BNC | Brysbaert-Biemiller AoA",
            "source_input_sha256": sha256(args.input), "morpholex_sha256": sha256(args.morpholex),
        })

    args.output.parent.mkdir(parents=True, exist_ok=True)
    fields = list(output_rows[0])
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader(); writer.writerows(output_rows)
    report = {
        "inputRows": len(output_rows), "existingCanonicalRows": sum(bool(row["existing_canonical_word_key"]) for row in output_rows),
        "newCandidateRows": sum(not row["existing_canonical_word_key"] for row in output_rows),
        "validTransformations": sum(row["spelling_transformation_status"] == "valid" for row in output_rows),
        "unresolvedTransformations": [row["display_word"] for row in output_rows if row["spelling_transformation_status"] != "valid"],
        "dictationsWithExactlyOneTarget": sum(row["provided_target_occurrences"] == "1" for row in output_rows),
        "stressedFinalSyllableCandidates": [row["display_word"] for row in output_rows if row["proposed_doubling_pattern"] == "stressed_final_syllable"],
        "requiresBritishFinalLRule": [row["display_word"] for row in output_rows if row["v1_lesson_eligibility"] == "requires_british_final_l_rule"],
        "lexicalCoverage": {
            "targetMorphoLex": sum(bool(row["target_morpholex_segmentation"]) for row in output_rows),
            "targetBritishIpa": sum(bool(row["target_british_ipa"]) for row in output_rows),
            "targetCmuDict": sum(bool(row["target_cmudict_phones"]) for row in output_rows),
            "targetAoA": sum(bool(row["target_aoa_median"]) for row in output_rows),
        },
        "missingBritishIpa": [row["display_word"] for row in output_rows if not row["target_british_ipa"]],
    }
    args.output.with_suffix(".report.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))

if __name__ == "__main__":
    main()
