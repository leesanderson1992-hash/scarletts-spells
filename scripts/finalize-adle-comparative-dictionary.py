#!/usr/bin/env python3
"""Project 25 approved comparative-degree words into the canonical release shape.

The evidence-only MorphoLex segmentation/POS columns are deliberately not
imported. The reviewed spelling analysis and child-facing word sum are the
project's own release facts; source evidence remains in the signed workbook.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook

SOURCE_KEY = "comparative_degree_katie_approved_2026_09_29"
SOURCE_NAME = "Katie Sanderson approved comparative-degree dictionary review"
SOURCE_URL = "local:comparative-dictionary-approved.xlsx"
SOURCE_LICENCE = "internal/project-reviewed factual compilation"


def rows_from_sheet(workbook: object, name: str) -> list[dict[str, str]]:
    sheet = workbook[name if name in workbook else name[:31]]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(values)]
    return [dict(zip(headers, (str(value).strip() if value is not None else "" for value in row))) for row in values if any(value is not None for value in row)]


def require_approved(rows: list[dict[str, str]], gates: list[str], label: str, expected: int) -> None:
    if len(rows) != expected:
        raise ValueError(f"{label}: expected {expected} rows, found {len(rows)}")
    for index, row in enumerate(rows, 2):
        missing = [gate for gate in gates if row.get(gate) != "approved"]
        if row.get("reviewed_by") != "Katie Sanderson" or row.get("reviewed_at") != "2026-09-29":
            missing.append("named_review")
        if missing:
            raise ValueError(f"{label} row {index}: unresolved {', '.join(missing)}")


def write_csv(path: Path, columns: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    selection = rows_from_sheet(workbook, "Canonical word review")
    morphology = rows_from_sheet(workbook, "Linguistic morphology & word sums")
    dictations = rows_from_sheet(workbook, "Dictation review")
    sources = rows_from_sheet(workbook, "Sources & licence")
    require_approved(selection, ["identity_review", "selection_evidence_review", "pronunciation_review", "british_english_review", "accessibility_review", "source_licence_review", "final_decision"], "comparative canonical words", 25)
    require_approved(morphology, ["linguistic_analysis_review", "word_sum_review", "final_decision"], "comparative morphology", 25)
    require_approved(dictations, ["child_language_review", "british_english_review", "accessibility_review", "final_decision"], "comparative dictations", 25)
    require_approved(sources, ["source_review"], "comparative evidence sources", 6)
    keys = [row["word_key"] for row in selection]
    if len(set(keys)) != 25 or set(keys) != {row["word_key"] for row in morphology} or set(keys) != {row["word_key"] for row in dictations}:
        raise ValueError("The three approved word sheets must contain the same 25 unique keys.")
    by_morphology = {row["word_key"]: row for row in morphology}
    by_dictation = {row["word_key"]: row for row in dictations}
    digest = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    note = f"Katie Sanderson reviewed 2026-09-29; workbook SHA-256 {digest}. External candidate evidence is recorded in the workbook; raw MorphoLex fields are not imported."
    provenance = {"source_category": "internal_reviewed_seed", "source_name": SOURCE_NAME, "source_url": SOURCE_URL, "source_licence": SOURCE_LICENCE, "source_use_note": note}
    approval = {"confidence": "high", "review_status": "approved_for_first_exposure"}
    words: list[dict[str, str]] = []
    metadata: list[dict[str, str]] = []
    morphs: list[dict[str, str]] = []
    sentences: list[dict[str, str]] = []
    for row in selection:
        key = row["word_key"]
        word = row["word"]
        m = by_morphology[key]
        d = by_dictation[key]
        if d["display_word"] != word or m["word"] != word:
            raise ValueError(f"{key}: approved word differs across sheets")
        if not all(row.get(field) for field in ("frequency_band", "age_band", "complexity_band", "british_ipa", "syllables", "stress_pattern", "has_schwa")):
            raise ValueError(f"{key}: incomplete approved word facts")
        if not all(m.get(field) for field in ("morphology_parts", "feature_keys", "morphology_joins", "word_sum")) or m["analysis_status"] != "approved":
            raise ValueError(f"{key}: incomplete approved morphology")
        if not all(d.get(field) for field in ("dictation_sentence", "dictation_target_token_index", "audio_text")):
            raise ValueError(f"{key}: incomplete approved dictation")
        words.append({"word_key": key, "normalised_word": word.lower(), "display_word": word, "dialect_code": "en-GB", "frequency_band": row["frequency_band"], "age_band": row["age_band"], "complexity_band": row["complexity_band"], **provenance, **approval, "row_status": "active"})
        metadata.append({"word_key": key, "syllables": row["syllables"], "phoneme_hint": row["british_ipa"], "grapheme_notes": f"Approved {row['micro_skill_key']} adjective spelling.", "stress_pattern": row["stress_pattern"], "has_schwa": row["has_schwa"], "morphemes": m["word_sum"], "morphology_notes": m["transformation_notes"], "irregularity_notes": "", **provenance, **approval})
        morphs.append({"word_key": key, "raw_morpholex_segmentation": "", "raw_morpholex_pos": "", "morphology_parts": m["morphology_parts"], "feature_keys": m["feature_keys"], "morphology_joins": m["morphology_joins"], "transformation_notes": m["transformation_notes"], "word_sum": m["word_sum"], "analysis_status": "approved", **provenance, **approval, "reviewed_by": m["reviewed_by"], "reviewed_at": m["reviewed_at"], "review_notes": "Reviewer-approved rule projection; external raw morphology retained only as workbook evidence."})
        sentences.append({"word_key": key, "display_word": word, "age_band": d["age_band"], "complexity_band": d["complexity_band"], "dictation_sentence": d["dictation_sentence"], "dictation_target_token_index": d["dictation_target_token_index"], "audio_text": d["audio_text"], **provenance, **approval, "reviewed_by": d["reviewed_by"], "reviewed_at": d["reviewed_at"], "review_notes": "Katie Sanderson approved 2026-09-29."})
    source = {"source_key": SOURCE_KEY, **provenance, "importability_status": "importable", "legal_review_status": "not_required"}
    args.output.mkdir(parents=True, exist_ok=True)
    write_csv(args.output / "canonical_words.csv", list(words[0]), words)
    write_csv(args.output / "canonical_word_metadata.csv", list(metadata[0]), metadata)
    write_csv(args.output / "canonical_word_morphology.csv", list(morphs[0]), morphs)
    write_csv(args.output / "dictation_sentences.csv", list(sentences[0]), sentences)
    write_csv(args.output / "teaching_content_sources.csv", list(source), [source])
    print(f"approved_comparative_degree_csv_emitted words={len(words)}")


if __name__ == "__main__":
    main()
